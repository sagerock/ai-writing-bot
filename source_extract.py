"""Text extraction with stable page or paragraph offset maps.

This module deliberately has no dependency on the web application so extraction
can be tested without importing FastAPI, Firebase, or model-provider SDKs.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from typing import Literal, TypedDict

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pypdf import PdfReader


class SourceMapEntry(TypedDict, total=False):
    page: int
    paragraph: int
    start: int
    end: int


class ExtractionResult(TypedDict):
    text: str
    pages: list[SourceMapEntry]
    kind: Literal["page", "paragraph"]


PDF_CONTENT_TYPES = frozenset({"application/pdf"})
DOCX_CONTENT_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/docx",
    }
)
XLSX_CONTENT_TYPES = frozenset(
    {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
)
# A 10 MB .xlsx can expand to millions of cells; stop reading well before that.
MAX_XLSX_CELLS = 500_000
CSV_CONTENT_TYPES = frozenset(
    {
        "text/csv",
        "application/csv",
        "application/vnd.ms-excel",
    }
)


def _normalized_content_type(content_type: str) -> str:
    return content_type.partition(";")[0].strip().lower()


def _decode_text(data: bytes) -> str:
    """Decode browser-uploaded text while removing an optional UTF-8 BOM."""
    return data.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")


def _paragraph_map(text: str) -> list[SourceMapEntry]:
    """Map non-empty blocks separated by one or more blank lines."""
    entries: list[SourceMapEntry] = []
    block_start = 0

    for separator in re.finditer(r"\n[ \t]*\n+", text):
        _append_paragraph_entry(entries, text, block_start, separator.start())
        block_start = separator.end()

    _append_paragraph_entry(entries, text, block_start, len(text))
    return entries


def _append_paragraph_entry(
    entries: list[SourceMapEntry], text: str, start: int, end: int
) -> None:
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    if start < end:
        entries.append(
            {
                "paragraph": len(entries) + 1,
                "start": start,
                "end": end,
            }
        )


def _join_paragraph_blocks(
    blocks: list[str], separator: str = "\n\n"
) -> tuple[str, list[SourceMapEntry]]:
    text_parts: list[str] = []
    entries: list[SourceMapEntry] = []
    offset = 0

    for block in (block.strip() for block in blocks):
        if not block:
            continue
        if text_parts:
            text_parts.append(separator)
            offset += len(separator)
        start = offset
        text_parts.append(block)
        offset += len(block)
        entries.append(
            {
                "paragraph": len(entries) + 1,
                "start": start,
                "end": offset,
            }
        )

    return "".join(text_parts), entries


def _extract_pdf(data: bytes) -> ExtractionResult:
    reader = PdfReader(io.BytesIO(data))
    page_texts = [page.extract_text() or "" for page in reader.pages]
    text_parts: list[str] = []
    pages: list[SourceMapEntry] = []
    offset = 0

    for page_number, page_text in enumerate(page_texts, start=1):
        if text_parts:
            text_parts.append("\n")
            offset += 1
        start = offset
        text_parts.append(page_text)
        offset += len(page_text)
        pages.append(
            {"page": page_number, "start": start, "end": offset}
        )

    return {"text": "".join(text_parts), "pages": pages, "kind": "page"}


def _escape_markdown_cell(value: str) -> str:
    return value.replace("|", r"\|").replace("\n", "<br>")


def _extract_csv(data: bytes) -> ExtractionResult:
    rows = list(csv.reader(io.StringIO(_decode_text(data))))
    if not rows:
        return {"text": "", "pages": [], "kind": "paragraph"}

    width = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in rows]

    def render_row(row: list[str]) -> str:
        return "| " + " | ".join(_escape_markdown_cell(cell) for cell in row) + " |"

    header = render_row(normalized_rows[0])
    divider = "| " + " | ".join(["---"] * width) + " |"
    blocks = [f"{header}\n{divider}"]
    blocks.extend(render_row(row) for row in normalized_rows[1:])
    text, pages = _join_paragraph_blocks(blocks, separator="\n")
    return {"text": text, "pages": pages, "kind": "paragraph"}


def _render_table(rows: list[list[str]]) -> str:
    width = max(len(row) for row in rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in rows]

    def render_row(row: list[str]) -> str:
        return "| " + " | ".join(_escape_markdown_cell(cell) for cell in row) + " |"

    divider = "| " + " | ".join(["---"] * width) + " |"
    return "\n".join(
        [render_row(normalized_rows[0]), divider]
        + [render_row(row) for row in normalized_rows[1:]]
    )


def _xlsx_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _extract_xlsx(data: bytes) -> ExtractionResult:
    """Render every non-empty worksheet as a titled Markdown table.

    Formula cells use the value Excel last calculated and saved.
    """
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[str] = []
    cells_read = 0
    truncated = False

    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for values in sheet.iter_rows(values_only=True):
                cells_read += len(values)
                if cells_read > MAX_XLSX_CELLS:
                    truncated = True
                    break
                row = [_xlsx_cell_text(value) for value in values]
                while row and not row[-1]:
                    row.pop()
                if row:
                    rows.append(row)
            if rows:
                blocks.append(f"## Sheet: {sheet.title}\n\n{_render_table(rows)}")
            if truncated:
                break
    finally:
        workbook.close()

    if truncated:
        blocks.append("[... spreadsheet truncated: too many cells to read ...]")
    text, pages = _join_paragraph_blocks(blocks)
    return {"text": text, "pages": pages, "kind": "paragraph"}


def _render_docx_table(table: Table) -> str:
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    rendered = ["\t".join(cell.strip() for cell in row) for row in rows]
    return "\n".join(rendered)


def _extract_docx(data: bytes) -> ExtractionResult:
    document = Document(io.BytesIO(data))
    blocks: list[str] = []

    # iter_inner_content preserves the order of paragraphs and tables.
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph):
            blocks.append(item.text)
        elif isinstance(item, Table):
            blocks.append(_render_docx_table(item))

    text, pages = _join_paragraph_blocks(blocks, separator="\n")
    return {"text": text, "pages": pages, "kind": "paragraph"}


def extract(data: bytes, content_type: str) -> ExtractionResult:
    """Extract source text and an offset map from a supported content type.

    ``pages`` contains ``page`` keys for PDFs and ``paragraph`` keys for all
    other formats. Every ``start``/``end`` pair is a slice into returned
    ``text``; numbering is one-based for use in citations.
    """
    normalized_type = _normalized_content_type(content_type)

    if normalized_type in PDF_CONTENT_TYPES:
        return _extract_pdf(data)
    if normalized_type in DOCX_CONTENT_TYPES:
        return _extract_docx(data)
    if normalized_type in XLSX_CONTENT_TYPES:
        return _extract_xlsx(data)
    if normalized_type in CSV_CONTENT_TYPES:
        return _extract_csv(data)
    if normalized_type.startswith("text/") or normalized_type in {
        "application/json",
        "application/xml",
        "application/sql",
        "application/x-sh",
        "application/x-yaml",
    }:
        text = _decode_text(data)
        return {"text": text, "pages": _paragraph_map(text), "kind": "paragraph"}

    raise ValueError(f"Unsupported content type: {content_type or '(missing)'}")
